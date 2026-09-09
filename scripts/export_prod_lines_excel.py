"""生产线回测面板 (PL 20 种子 + SEK 生产 5 种子集成点) -> 一份 Excel

对应活文档 docs/prod_lines_backtest_latest.md: 每次 te 变了在 040 重跑 PL/SEK 后,
把 data/processed/wf_daily_PL_*.json / wf_daily_SEK_*_k5_0_*.json 同步到本地再跑本脚本.

默认输出到桌面: ~/Desktop/l量化操作表/<生成日期>/生产线回测_te<te>.xlsx

用法:
  python scripts/export_prod_lines_excel.py                # 自动找最新 te
  python scripts/export_prod_lines_excel.py --te 2026-09-04 --output /路径.xlsx
"""
import argparse
import glob
import json
import re
import statistics as st
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
PROC = ROOT / "data" / "processed"
DESKTOP_ROOT = Path.home() / "Desktop" / "l量化操作表"

# 顺序与活文档一致; 描述 = live_config 逐字配置的人话版
LINES = [("aggr5w", "5万/3只/主板/T1A/ind-cap 2"),
         ("aggr10w", "10万/3只/主板/T1A/ind-cap 2"),
         ("steady5w", "5万/5只/主板/T1B"),
         ("fyf100w", "100万/8只/主板/T1B"),
         ("aggr2w", "2万/2只/主板/基线80/lot-flex 0.5"),
         ("steady2w", "2万/3只/全市场/基线80/lot-flex 0.5/ind-cap 2")]

NOTES = [
    "窗 2023-09-19 → te (约 3 年), 生产矩阵 training_data_pit_v24_tick1, 每条线用 live_config.signal_args 的逐字配置.",
    "「逐种子」表是把生产的 --seed-ensemble 拆成 20 个单种子看分布; 生产真跑的是 5 种子集成 (42/7/123/2024/31337),",
    "  它的逐字等价点在「六线汇总」最后几列和「逐种子」里 种子=生产5种子集成 那几行. 报中位, 不报单次.",
    "样本内性质: min-pred / roll-rank / fill-daily / ind-cap / T1A·T1B 都是在 2022-09~2026-07 上挑出来的, 这段窗对现行配置是样本内;",
    "  只有上次 te 之后的那一段是真样本外. 它回答『现行配置到今天长什么样』, 不回答『能不能持续』.",
    "弱窗风险账 (findings_2026-09-05 §): 2020-07~2022-08 生产配方 20 种子 -36.8%, 0/20 正, 0/20 赢基准; 没有已验证的防御, 只能靠仓位/本金上限.",
    "生产种子组在 aggr5w/aggr10w 上恰好是 12 个 k=5 子集里最差的, 是运气不是缺陷; 不拿回测挑种子 (SEK 09-07).",
    "基准 = 等权买入持有: 训练线(aggr5w/steady5w/aggr2w)是主板等权, 重放线(aggr10w/fyf100w)与全市场线是全市场等权; 只影响 基准/IR 两列.",
    "评价策略看 IC 及其 t 值, 不要看总收益 (findings_2026-07-29): 只持 3~5 只时总收益的 t 值约 1, 统计上无法与 0 区分.",
]


def latest_te():
    tes = set()
    for p in PROC.glob("wf_daily_PL_*_te*.json"):
        m = re.search(r"_te(\d{4}-\d{2}-\d{2})_", p.name)
        if m:
            tes.add(m.group(1))
    if not tes:
        raise SystemExit(f"{PROC} 下没有 wf_daily_PL_*.json, 先从 040 同步")
    return max(tes)


def load(path: Path):
    d = json.loads(path.read_text(encoding="utf-8"))
    s = d["summary"]
    rec = {r["segment"]: r for r in d.get("recent", [])}
    stab = {r["segment"]: r for r in d.get("stability", [])}
    row = {
        "总收益%": s["total_return_pct"], "年化%": s["annualized_return_pct"],
        "夏普": s["sharpe"], "最大回撤%": s["max_dd_pct"], "IR": s["information_ratio"],
        "基准总收益%": s["benchmark_total_pct"], "赢基准": s["total_return_pct"] > s["benchmark_total_pct"],
        "超额年化%": s["excess_annual_pct"], "IC均值": s["ic_mean"], "IC_t": s["ic_tstat"],
        "总费用%本金": s["total_cost_pct"], "空仓天数%": s["cash_days_pct"],
        "平均仓位%": s["avg_deployed_pct"], "平均持仓只数": s["avg_holdings"],
        "交易笔数": s["n_trades"], "续持次数": s.get("n_rolled"), "日补买次数": s.get("n_daily_fill"),
        "低于门槛不买": s.get("n_below_thresh"), "行业上限跳过": s.get("n_ind_skip"),
        "期末净值": s["final_value"], "本金": d["initial_capital"],
        "前半段策略%": stab.get("前半段", {}).get("strategy_pct"),
        "前半段基准%": stab.get("前半段", {}).get("benchmark_pct"),
        "后半段策略%": stab.get("后半段", {}).get("strategy_pct"),
        "后半段基准%": stab.get("后半段", {}).get("benchmark_pct"),
        "最近6月策略%": rec.get("最近6月", {}).get("strategy_pct"),
        "最近6月基准%": rec.get("最近6月", {}).get("benchmark_pct"),
        "最近3月策略%": rec.get("最近3月", {}).get("strategy_pct"),
        "最近3月基准%": rec.get("最近3月", {}).get("benchmark_pct"),
        "区间": d["period"], "文件": path.name,
    }
    cfg = {k: d.get(k) for k in (
        "train_file", "pit_universe", "label", "objective", "hold_days", "tranche_n",
        "initial_capital", "portfolio_mode", "exec_mode", "slippage", "trade_cost", "min_fee",
        "regime_filter", "regime_ma", "regime_breadth", "regime_confirm", "min_pred",
        "fill_daily", "roll_rank", "lot_flex", "skip_boards", "features")}
    cfg["ind_cap"] = s.get("ind_cap")
    daily = pd.DataFrame(d["daily"])[["date", "portfolio_value"]]
    daily["净值"] = daily["portfolio_value"] / d["initial_capital"]
    return row, cfg, daily.set_index("date")["净值"]


def q(vals, fmt="{:+.1f}"):
    return fmt.format(st.median(vals)) if vals else "—"


def build(te: str, out: Path):
    seed_rows, cfgs, curves_prod, curves_seed = [], {}, {}, {}
    summary_rows = []
    for line, desc in LINES:
        files = sorted(PROC.glob(f"wf_daily_PL_{line}_s*_te{te}_*.json"))
        if not files:
            print(f"!! {line}: 没有 PL 结果, 跳过")
            continue
        rows = []
        for f in files:
            seed = int(re.search(r"_s(\d+)_ts", f.name).group(1))
            row, cfg, curve = load(f)
            row = {"线": line, "配置": desc, "种子": seed, **row}
            rows.append(row); seed_rows.append(row)
            curves_seed.setdefault(line, {})[seed] = curve
            cfgs.setdefault(line, {"线": line, "配置": desc, **cfg})
        prod = None
        pf = sorted(PROC.glob(f"wf_daily_SEK_{line}_k5_0_s42_*_te{te}_*.json"))
        if pf:
            prod, _, pcurve = load(pf[0])
            seed_rows.append({"线": line, "配置": desc, "种子": "生产5种子集成", **prod})
            curves_prod[line] = pcurve

        v = [r["总收益%"] for r in rows]
        qs = st.quantiles(v, n=4) if len(v) >= 2 else [v[0]] * 3
        r6 = [(r["最近6月策略%"], r["最近6月基准%"]) for r in rows if r["最近6月策略%"] is not None]
        r3 = [(r["最近3月策略%"], r["最近3月基准%"]) for r in rows if r["最近3月策略%"] is not None]
        summary_rows.append({
            "线": line, "配置": desc, "种子数": len(v),
            "总收益 中位%": round(st.median(v), 1), "Q1%": round(qs[0], 1), "Q3%": round(qs[2], 1),
            "最差%": min(v), "最好%": max(v),
            "正收益种子": f"{sum(x > 0 for x in v)}/{len(v)}",
            "赢基准种子": f"{sum(r['赢基准'] for r in rows)}/{len(v)}",
            "基准总收益%": rows[0]["基准总收益%"],
            "年化 中位%": q([r["年化%"] for r in rows]), "夏普 中位": q([r["夏普"] for r in rows], "{:.2f}"),
            "最大回撤 中位%": q([r["最大回撤%"] for r in rows]), "IR 中位": q([r["IR"] for r in rows], "{:+.2f}"),
            "IC_t 中位": q([r["IC_t"] for r in rows], "{:.2f}"),
            "费用 中位%": q([r["总费用%本金"] for r in rows], "{:.1f}"),
            "空仓 中位%": q([r["空仓天数%"] for r in rows], "{:.0f}"),
            "前半段 中位%": q([r["前半段策略%"] for r in rows if r["前半段策略%"] is not None]),
            "后半段 中位%": q([r["后半段策略%"] for r in rows if r["后半段策略%"] is not None]),
            "最近6月 策略中位/基准": f"{st.median([a for a, _ in r6]):+.1f} / {r6[0][1]:+.1f} ({sum(a > b for a, b in r6)}/{len(r6)} 赢)" if r6 else "—",
            "最近3月 策略中位/基准": f"{st.median([a for a, _ in r3]):+.1f} / {r3[0][1]:+.1f} ({sum(a > b for a, b in r3)}/{len(r3)} 赢)" if r3 else "—",
            "生产5种子 总收益%": prod["总收益%"] if prod else None,
            "生产5种子 夏普": prod["夏普"] if prod else None,
            "生产5种子 回撤%": prod["最大回撤%"] if prod else None,
            "生产5种子 最近6月%": prod["最近6月策略%"] if prod else None,
            "生产5种子 最近3月%": prod["最近3月策略%"] if prod else None,
        })

    summary = pd.DataFrame(summary_rows)
    seeds = pd.DataFrame(seed_rows)
    cfg_df = pd.DataFrame(list(cfgs.values()))
    cfg_df["skip_boards"] = cfg_df["skip_boards"].map(lambda x: ",".join(map(str, x)) if isinstance(x, list) else x)
    prod_nav = pd.DataFrame(curves_prod)
    band = {}
    for line, cs in curves_seed.items():
        df = pd.DataFrame(cs)
        band[f"{line} 最差"] = df.min(axis=1)
        band[f"{line} 中位"] = df.median(axis=1)
        band[f"{line} 最好"] = df.max(axis=1)
    band_nav = pd.DataFrame(band)
    notes = pd.DataFrame({"说明": [f"te = {te}; 生成于 {datetime.now():%Y-%m-%d %H:%M}; 来源 data/processed/wf_daily_PL_*.json + wf_daily_SEK_*_k5_0_*.json"] + NOTES})

    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="六线汇总", index=False)
        seeds.to_excel(xw, sheet_name="逐种子", index=False)
        prod_nav.round(4).to_excel(xw, sheet_name="净值_生产5种子", index_label="日期")
        band_nav.round(4).to_excel(xw, sheet_name="净值_20种子带", index_label="日期")
        cfg_df.to_excel(xw, sheet_name="配置", index=False)
        notes.to_excel(xw, sheet_name="说明", index=False)
    format_workbook(out)
    return summary


def format_workbook(path: Path):
    book = load_workbook(path)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(color="FFFFFF", bold=True)
    for ws in book.worksheets:
        ws.freeze_panes = "B2" if ws.title.startswith("净值") else "A2"
        if ws.title in ("逐种子",):
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for col in ws.columns:
            w = max(len(str(c.value or "")) for c in col[:60]) + 2
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w, 9), 60 if ws.title == "说明" else 30)
        hdr = {c.value: c.column for c in ws[1]}
        for key in ("总收益%", "总收益 中位%", "最近6月策略%", "最近3月策略%", "生产5种子 总收益%"):
            if key in hdr and ws.max_row > 1:
                L = get_column_letter(hdr[key])
                ws.conditional_formatting.add(
                    f"{L}2:{L}{ws.max_row}",
                    ColorScaleRule(start_type="min", start_color="F4CCCC",
                                   mid_type="num", mid_value=0, mid_color="FFFFFF",
                                   end_type="max", end_color="D9EAD3"))
    book.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--te", default=None, help="测试终点, 缺省取 data/processed 里最新的")
    ap.add_argument("--output", type=Path, default=None)
    a = ap.parse_args()
    te = a.te or latest_te()
    out = a.output or DESKTOP_ROOT / datetime.now().strftime("%Y-%m-%d") / f"生产线回测_te{te}.xlsx"
    summary = build(te, out)
    with pd.option_context("display.width", 250, "display.max_columns", 40):
        print(summary[["线", "总收益 中位%", "Q1%", "Q3%", "最差%", "正收益种子", "赢基准种子",
                       "夏普 中位", "最大回撤 中位%", "前半段 中位%", "后半段 中位%",
                       "生产5种子 总收益%"]].to_string(index=False))
    print(f"\n已写入 {out}")


if __name__ == "__main__":
    main()
